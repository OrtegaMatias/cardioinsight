"""Exporta el HTML a PDF vectorial y a un respaldo PPTX de 3840 × 2160."""

import subprocess
from pathlib import Path
from tempfile import TemporaryDirectory

from pptx import Presentation
from pptx.util import Inches

ROOT = Path(__file__).resolve().parent
HTML = ROOT / "PRESENTACION_INTERACTIVA.html"
PDF = ROOT / "PRESENTACION_CERTAMEN.pdf"
ASSETS = ROOT / "presentacion_web_assets"
CHROME = Path("/Applications/Google Chrome.app/Contents/MacOS/Google Chrome")

if not CHROME.is_file():
    raise SystemExit(f"Chrome no está instalado en {CHROME}")
if not HTML.is_file():
    raise SystemExit(f"Falta la presentación {HTML}")
ASSETS.mkdir(exist_ok=True)

with TemporaryDirectory(prefix="certamen-slides-") as profile:
    subprocess.run(
        [
            str(CHROME),
            "--headless",
            "--disable-gpu",
            "--no-first-run",
            "--no-pdf-header-footer",
            f"--user-data-dir={profile}",
            f"--print-to-pdf={PDF}",
            "--window-size=1600,900",
            "--virtual-time-budget=3000",
            HTML.as_uri(),
        ],
        check=True,
        stdout=subprocess.DEVNULL,
        stderr=subprocess.PIPE,
        text=True,
    )

pdf_info = subprocess.check_output(["pdfinfo", str(PDF)], text=True)
page_count = int(next(line.split(":", 1)[1] for line in pdf_info.splitlines() if line.startswith("Pages:")))
if page_count != 12:
    raise SystemExit(f"La impresión produjo {page_count} páginas; se esperaban 12. Revisar CSS de impresión.")

presentation = Presentation()
presentation.slide_width = Inches(13.333333)
presentation.slide_height = Inches(7.5)
blank = presentation.slide_layouts[6]
for number in range(1, page_count + 1):
    image = ASSETS / f"slide_{number:02d}.png"
    subprocess.run(
        [
            "pdftoppm",
            "-f", str(number),
            "-l", str(number),
            "-scale-to-x", "3840",
            "-scale-to-y", "2160",
            "-png",
            "-singlefile",
            str(PDF),
            str(image.with_suffix("")),
        ],
        check=True,
        stdout=subprocess.DEVNULL,
    )
    slide = presentation.slides.add_slide(blank)
    slide.shapes.add_picture(
        str(image),
        0, 0,
        width=presentation.slide_width,
        height=presentation.slide_height,
    )

presentation.save(ROOT / "PRESENTACION_CERTAMEN.pptx")
print(f"PDF vectorial y PPTX 4K creados a partir del HTML; {page_count} láminas 16:9.")
